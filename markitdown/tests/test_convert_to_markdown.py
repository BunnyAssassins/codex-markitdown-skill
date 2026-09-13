from __future__ import annotations

import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document
from openpyxl import Workbook
from pptx import Presentation


SKILL_DIR = Path(__file__).resolve().parents[1]
SCRIPT = SKILL_DIR / "scripts" / "convert_to_markdown.py"


class ConversionTests(unittest.TestCase):
    def run_converter(self, *arguments: object) -> subprocess.CompletedProcess[str]:
        return subprocess.run(
            [sys.executable, str(SCRIPT), *map(str, arguments)],
            text=True,
            capture_output=True,
            check=False,
        )

    def test_office_and_text_batch(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            inputs = root / "inputs"
            outputs = root / "outputs"
            inputs.mkdir()

            text_file = inputs / "notes.txt"
            text_file.write_text("Codex MarkItDown smoke test", encoding="utf-8")

            docx_file = inputs / "brief.docx"
            document = Document()
            document.add_heading("Quarterly Brief", level=1)
            document.add_paragraph("Revenue increased in the test fixture.")
            document.save(docx_file)

            xlsx_file = inputs / "metrics.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Metrics"
            sheet.append(["Metric", "Value"])
            sheet.append(["Accuracy", 0.91])
            workbook.save(xlsx_file)

            pptx_file = inputs / "slides.pptx"
            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[1])
            slide.shapes.title.text = "Launch Plan"
            slide.placeholders[1].text = "Verify the converted slide text."
            presentation.save(pptx_file)

            result = self.run_converter(
                text_file, docx_file, xlsx_file, pptx_file, "--output-dir", outputs
            )
            self.assertEqual(result.returncode, 0, result.stderr)
            self.assertIn("Codex MarkItDown smoke test", (outputs / "notes.md").read_text(encoding="utf-8"))
            self.assertIn("Quarterly Brief", (outputs / "brief.md").read_text(encoding="utf-8"))
            self.assertIn("Accuracy", (outputs / "metrics.md").read_text(encoding="utf-8"))
            self.assertIn("Launch Plan", (outputs / "slides.md").read_text(encoding="utf-8"))

    def test_refuses_overwrite(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "sample.txt"
            output = root / "output"
            source.write_text("first", encoding="utf-8")
            output.mkdir()
            (output / "sample.md").write_text("keep me", encoding="utf-8")

            result = self.run_converter(source, "--output-dir", output)
            self.assertEqual(result.returncode, 1)
            self.assertEqual((output / "sample.md").read_text(encoding="utf-8"), "keep me")

    def test_stdout_and_url_rejection(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "sample.txt"
            source.write_text("stdout content", encoding="utf-8")
            result = self.run_converter(source, "--stdout")
            self.assertEqual(result.returncode, 0, result.stderr)
            self.assertIn("stdout content", result.stdout)

        rejected = self.run_converter("https://example.com/file.pdf", "--stdout")
        self.assertEqual(rejected.returncode, 2)
        self.assertIn("remote URLs are not accepted", rejected.stderr)


if __name__ == "__main__":
    unittest.main()

